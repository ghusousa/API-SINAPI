"""Carregador de dados SINAPI a partir de arquivos XLSX/ZIP da Caixa.

Processa os arquivos mensais SINAPI no formato XLSX (baixados de
https://www.caixa.gov.br/sinapi) e carrega os dados de insumos,
composições e composições analíticas em memória.

O formato esperado segue o padrão oficial da Caixa Econômica Federal:
- Planilhas de Insumos: colunas CODIGO, DESCRICAO DO INSUMO, UNIDADE, PRECO MEDIANO
- Planilhas de Composições: colunas CODIGO, DESCRICAO DA COMPOSICAO, UNIDADE, CUSTO TOTAL
- Planilhas Analíticas: detalhamento dos itens que compõem cada composição
"""

import io
import os
import re
import zipfile
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


# Mapeamento UF -> Nome do estado (usado para identificar arquivos no ZIP)
UF_NAMES = {
    "AC": "Acre", "AL": "Alagoas", "AM": "Amazonas", "AP": "Amapa",
    "BA": "Bahia", "CE": "Ceara", "DF": "Distrito Federal",
    "ES": "Espirito Santo", "GO": "Goias", "MA": "Maranhao",
    "MG": "Minas Gerais", "MS": "Mato Grosso do Sul", "MT": "Mato Grosso",
    "PA": "Para", "PB": "Paraiba", "PE": "Pernambuco", "PI": "Piaui",
    "PR": "Parana", "RJ": "Rio de Janeiro", "RN": "Rio Grande do Norte",
    "RO": "Rondonia", "RR": "Roraima", "RS": "Rio Grande do Sul",
    "SC": "Santa Catarina", "SE": "Sergipe", "SP": "Sao Paulo",
    "TO": "Tocantins",
}

ESTADOS_INFO = {
    "AC": {"nome": "Acre", "ibge": 12, "regiao": "norte"},
    "AL": {"nome": "Alagoas", "ibge": 27, "regiao": "nordeste"},
    "AM": {"nome": "Amazonas", "ibge": 13, "regiao": "norte"},
    "AP": {"nome": "Amapá", "ibge": 16, "regiao": "norte"},
    "BA": {"nome": "Bahia", "ibge": 29, "regiao": "nordeste"},
    "CE": {"nome": "Ceará", "ibge": 23, "regiao": "nordeste"},
    "DF": {"nome": "Distrito Federal", "ibge": 53, "regiao": "centro-oeste"},
    "ES": {"nome": "Espírito Santo", "ibge": 32, "regiao": "sudeste"},
    "GO": {"nome": "Goiás", "ibge": 52, "regiao": "centro-oeste"},
    "MA": {"nome": "Maranhão", "ibge": 21, "regiao": "nordeste"},
    "MG": {"nome": "Minas Gerais", "ibge": 31, "regiao": "sudeste"},
    "MS": {"nome": "Mato Grosso do Sul", "ibge": 50, "regiao": "centro-oeste"},
    "MT": {"nome": "Mato Grosso", "ibge": 51, "regiao": "centro-oeste"},
    "PA": {"nome": "Pará", "ibge": 15, "regiao": "norte"},
    "PB": {"nome": "Paraíba", "ibge": 25, "regiao": "nordeste"},
    "PE": {"nome": "Pernambuco", "ibge": 26, "regiao": "nordeste"},
    "PI": {"nome": "Piauí", "ibge": 22, "regiao": "nordeste"},
    "PR": {"nome": "Paraná", "ibge": 41, "regiao": "sul"},
    "RJ": {"nome": "Rio de Janeiro", "ibge": 33, "regiao": "sudeste"},
    "RN": {"nome": "Rio Grande do Norte", "ibge": 24, "regiao": "nordeste"},
    "RO": {"nome": "Rondônia", "ibge": 11, "regiao": "norte"},
    "RR": {"nome": "Roraima", "ibge": 14, "regiao": "norte"},
    "RS": {"nome": "Rio Grande do Sul", "ibge": 43, "regiao": "sul"},
    "SC": {"nome": "Santa Catarina", "ibge": 42, "regiao": "sul"},
    "SE": {"nome": "Sergipe", "ibge": 28, "regiao": "nordeste"},
    "SP": {"nome": "São Paulo", "ibge": 35, "regiao": "sudeste"},
    "TO": {"nome": "Tocantins", "ibge": 17, "regiao": "norte"},
}


def _normalise(text: str) -> str:
    """Remove acentos e converte para maiúsculas para comparação."""
    import unicodedata
    nfkd = unicodedata.normalize("NFKD", text)
    return "".join(c for c in nfkd if not unicodedata.combining(c)).upper()


def _detect_regime(name_or_path: str) -> str:
    """Detecta o regime (DESONERADO/NAO_DESONERADO) pelo nome da aba ou arquivo."""
    name = _normalise(name_or_path)
    if "NAODESON" in name.replace(" ", "").replace("_", ""):
        return "NAO_DESONERADO"
    if "NAO" in name and "DESON" in name:
        return "NAO_DESONERADO"
    if "SEM" in name and "DESON" in name:
        return "NAO_DESONERADO"
    if "DESON" in name:
        return "DESONERADO"
    # Siglas padrão SINAPI
    if name.startswith("ISD") or name.startswith("CSD"):
        return "NAO_DESONERADO"
    if name.startswith("ICD") or name.startswith("CCD"):
        return "DESONERADO"
    return "NAO_DESONERADO"


def _detect_sheet_type(sheet_name: str) -> Optional[str]:
    """Detecta se a aba é de insumos, composições ou analítico."""
    name = _normalise(sheet_name)
    if "ANALITICO" in name or "ANALITICA" in name:
        return "analitico"
    if "INSUMO" in name or name.startswith("IS") or name.startswith("IC"):
        return "insumo"
    if "COMPOS" in name or name.startswith("CS") or name.startswith("CC"):
        return "composicao"
    return None


def _detect_file_type_from_filename(filename: str) -> Optional[str]:
    """Detecta tipo de dados pelo padrão de nome do arquivo XLSX real SINAPI."""
    name = _normalise(os.path.basename(filename))
    if "ANALITICO" in name or "ANALITICA" in name:
        return "analitico"
    if "SINTETICO" in name or "SINTETICA" in name:
        return "composicao"
    if "PRECO" in name and "INSUMO" in name:
        return "insumo"
    return None


def _extract_hyperlink_value(val):
    """Extrai valor numérico de fórmulas HYPERLINK em células CODIGO.

    Células SINAPI podem conter fórmulas como =HYPERLINK("...", 12345).
    """
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return val
    s = str(val).strip()
    m = re.match(r'=HYPERLINK\("(?:[^"\\]|\\.)*",\s*"?(\d+)"?\)', s, re.IGNORECASE)
    if m:
        return int(m.group(1))
    return val


def _find_header_row(ws, max_rows=50) -> Tuple[Optional[int], Dict[str, int]]:
    """Encontra a linha de cabeçalho e mapeia colunas por nome."""
    for row_idx in range(1, max_rows + 1):
        cells = {
            _normalise(str(c.value or "")): c.column - 1
            for c in ws[row_idx]
            if c.value
        }
        # Procura pela coluna CODIGO que é comum em todas as abas
        for key in cells:
            if "CODIGO" in key:
                return row_idx, cells
    return None, {}


def _safe_float(val) -> Optional[float]:
    """Converte valor para float de forma segura."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", ".").strip())
    except (ValueError, TypeError):
        return None


def _safe_int(val) -> Optional[int]:
    """Converte valor para int de forma segura, tratando fórmulas HYPERLINK."""
    if val is None:
        return None
    val = _extract_hyperlink_value(val)
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return None


def _find_col(columns: Dict[str, int], *keywords) -> Optional[int]:
    """Encontra índice de coluna que contém qualquer uma das keywords."""
    for col_name, idx in columns.items():
        for kw in keywords:
            if kw in col_name:
                return idx
    return None


def parse_xlsx_workbook(wb, estado: str, referencia: str) -> dict:
    """Processa um workbook XLSX e retorna dados estruturados.

    Args:
        wb: Workbook openpyxl.
        estado: UF (ex: 'SP').
        referencia: Data de referência (ex: '2026-01').

    Returns:
        dict com chaves 'insumos', 'composicoes', 'analitico'.
    """
    result = {"insumos": [], "composicoes": [], "analitico": []}
    estado = estado.upper()

    for sheet_name in wb.sheetnames:
        sheet_type = _detect_sheet_type(sheet_name)
        if sheet_type is None:
            continue

        regime = _detect_regime(sheet_name)
        ws = wb[sheet_name]
        header_row, columns = _find_header_row(ws)
        if header_row is None:
            continue

        if sheet_type == "insumo":
            col_codigo = _find_col(columns, "CODIGO")
            col_desc = _find_col(columns, "DESCRICAO")
            col_unidade = _find_col(columns, "UNIDADE")
            col_preco = _find_col(columns, "PRECO", "MEDIANO", "CUSTO")

            if col_codigo is None or col_desc is None:
                continue

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                codigo = _safe_int(row[col_codigo] if col_codigo < len(row) else None)
                if not codigo:
                    continue
                desc = str(row[col_desc] if col_desc < len(row) else "") or ""
                unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
                preco = _safe_float(row[col_preco] if col_preco is not None and col_preco < len(row) else None)

                result["insumos"].append({
                    "codigo": codigo,
                    "descricao": desc.strip(),
                    "unidade": unidade.strip(),
                    "preco_mediano": preco,
                    "estado": estado,
                    "regime": regime,
                    "referencia": referencia,
                    "fonte": "SINAPI",
                })

        elif sheet_type == "composicao":
            col_codigo = _find_col(columns, "CODIGO")
            col_desc = _find_col(columns, "DESCRICAO")
            col_unidade = _find_col(columns, "UNIDADE")
            col_custo = _find_col(columns, "CUSTO", "TOTAL", "PRECO")

            if col_codigo is None or col_desc is None:
                continue

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                codigo = _safe_int(row[col_codigo] if col_codigo < len(row) else None)
                if not codigo:
                    continue
                desc = str(row[col_desc] if col_desc < len(row) else "") or ""
                unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
                custo = _safe_float(row[col_custo] if col_custo is not None and col_custo < len(row) else None)

                result["composicoes"].append({
                    "codigo": codigo,
                    "descricao": desc.strip(),
                    "unidade": unidade.strip(),
                    "custo_total": custo,
                    "estado": estado,
                    "regime": regime,
                    "referencia": referencia,
                    "fonte": "SINAPI",
                })

        elif sheet_type == "analitico":
            col_comp = _find_col(columns, "COMPOSICAO", "COMP")
            col_item = _find_col(columns, "ITEM", "INSUMO", "CODIGO")
            col_tipo = _find_col(columns, "TIPO")
            col_desc = _find_col(columns, "DESCRICAO")
            col_unidade = _find_col(columns, "UNIDADE")
            col_coef = _find_col(columns, "COEFICIENTE", "QUANTIDADE", "QUANT")
            col_preco = _find_col(columns, "PRECO", "CUSTO", "UNITARIO")

            # col_item pode ser o mesmo que col_comp em certas planilhas
            if col_comp is None:
                col_comp = col_item

            if col_item is None:
                continue

            current_comp = None
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                comp_val = _safe_int(row[col_comp] if col_comp is not None and col_comp < len(row) else None)
                item_val = _safe_int(row[col_item] if col_item < len(row) else None)

                if comp_val and not item_val:
                    current_comp = comp_val
                    continue

                if item_val and current_comp:
                    desc = str(row[col_desc] if col_desc is not None and col_desc < len(row) else "") or ""
                    unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
                    coef = _safe_float(row[col_coef] if col_coef is not None and col_coef < len(row) else None)
                    preco = _safe_float(row[col_preco] if col_preco is not None and col_preco < len(row) else None)

                    tipo = "INSUMO"
                    if col_tipo is not None and col_tipo < len(row) and row[col_tipo]:
                        t = _normalise(str(row[col_tipo]))
                        if "COMP" in t:
                            tipo = "COMPOSICAO"

                    result["analitico"].append({
                        "composicao_codigo": current_comp,
                        "item_codigo": item_val,
                        "tipo_item": tipo,
                        "descricao": desc.strip(),
                        "unidade": unidade.strip(),
                        "coeficiente": coef or 0.0,
                        "preco_unitario": preco,
                        "estado": estado,
                        "regime": regime,
                        "referencia": referencia,
                    })

    return result


def parse_single_type_xlsx(wb, file_type: str, estado: str, referencia: str,
                           regime: str) -> dict:
    """Processa um workbook XLSX de tipo único (formato real SINAPI com arquivos separados).

    No formato real da Caixa, cada arquivo XLSX contém apenas um tipo de dado
    (insumos, composições sintético ou analítico) em uma única aba de dados.

    Args:
        wb: Workbook openpyxl.
        file_type: 'insumo', 'composicao' ou 'analitico'.
        estado: UF (ex: 'SP').
        referencia: Data de referência (ex: '2026-01').
        regime: 'DESONERADO' ou 'NAO_DESONERADO'.

    Returns:
        dict com chaves 'insumos', 'composicoes', 'analitico'.
    """
    result = {"insumos": [], "composicoes": [], "analitico": []}
    estado = estado.upper()

    # Use the first (or active) sheet
    ws = wb.active or wb[wb.sheetnames[0]]
    header_row, columns = _find_header_row(ws)
    if header_row is None:
        return result

    if file_type == "insumo":
        col_codigo = _find_col(columns, "CODIGO")
        col_desc = _find_col(columns, "DESCRICAO")
        col_unidade = _find_col(columns, "UNIDADE")
        col_preco = _find_col(columns, "PRECO", "MEDIANO", "CUSTO")

        if col_codigo is None or col_desc is None:
            return result

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            codigo = _safe_int(row[col_codigo] if col_codigo < len(row) else None)
            if not codigo:
                continue
            desc = str(row[col_desc] if col_desc < len(row) else "") or ""
            unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
            preco = _safe_float(row[col_preco] if col_preco is not None and col_preco < len(row) else None)

            result["insumos"].append({
                "codigo": codigo,
                "descricao": desc.strip(),
                "unidade": unidade.strip(),
                "preco_mediano": preco,
                "estado": estado,
                "regime": regime,
                "referencia": referencia,
                "fonte": "SINAPI",
            })

    elif file_type == "composicao":
        col_codigo = _find_col(columns, "CODIGO")
        col_desc = _find_col(columns, "DESCRICAO")
        col_unidade = _find_col(columns, "UNIDADE")
        col_custo = _find_col(columns, "CUSTO", "TOTAL", "PRECO")

        if col_codigo is None or col_desc is None:
            return result

        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            codigo = _safe_int(row[col_codigo] if col_codigo < len(row) else None)
            if not codigo:
                continue
            desc = str(row[col_desc] if col_desc < len(row) else "") or ""
            unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
            custo = _safe_float(row[col_custo] if col_custo is not None and col_custo < len(row) else None)

            result["composicoes"].append({
                "codigo": codigo,
                "descricao": desc.strip(),
                "unidade": unidade.strip(),
                "custo_total": custo,
                "estado": estado,
                "regime": regime,
                "referencia": referencia,
                "fonte": "SINAPI",
            })

    elif file_type == "analitico":
        col_comp = _find_col(columns, "COMPOSICAO", "COMP")
        col_item = _find_col(columns, "ITEM", "INSUMO", "CODIGO")
        col_tipo = _find_col(columns, "TIPO")
        col_desc = _find_col(columns, "DESCRICAO")
        col_unidade = _find_col(columns, "UNIDADE")
        col_coef = _find_col(columns, "COEFICIENTE", "QUANTIDADE", "QUANT")
        col_preco = _find_col(columns, "PRECO", "CUSTO", "UNITARIO")

        if col_comp is None:
            col_comp = col_item
        if col_item is None:
            return result

        current_comp = None
        for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
            comp_val = _safe_int(row[col_comp] if col_comp is not None and col_comp < len(row) else None)
            item_val = _safe_int(row[col_item] if col_item < len(row) else None)

            if comp_val and not item_val:
                current_comp = comp_val
                continue

            if item_val and current_comp:
                desc = str(row[col_desc] if col_desc is not None and col_desc < len(row) else "") or ""
                unidade = str(row[col_unidade] if col_unidade is not None and col_unidade < len(row) else "") or ""
                coef = _safe_float(row[col_coef] if col_coef is not None and col_coef < len(row) else None)
                preco = _safe_float(row[col_preco] if col_preco is not None and col_preco < len(row) else None)

                tipo = "INSUMO"
                if col_tipo is not None and col_tipo < len(row) and row[col_tipo]:
                    t = _normalise(str(row[col_tipo]))
                    if "COMP" in t:
                        tipo = "COMPOSICAO"

                result["analitico"].append({
                    "composicao_codigo": current_comp,
                    "item_codigo": item_val,
                    "tipo_item": tipo,
                    "descricao": desc.strip(),
                    "unidade": unidade.strip(),
                    "coeficiente": coef or 0.0,
                    "preco_unitario": preco,
                    "estado": estado,
                    "regime": regime,
                    "referencia": referencia,
                })

    return result


def load_xlsx_file(file_path_or_bytes, estado: str, referencia: str) -> dict:
    """Carrega dados de um arquivo XLSX SINAPI.

    Args:
        file_path_or_bytes: Caminho do arquivo ou bytes.
        estado: UF do estado.
        referencia: Data de referência (YYYY-MM).

    Returns:
        dict com dados estruturados.
    """
    if isinstance(file_path_or_bytes, (bytes, io.BytesIO)):
        buf = io.BytesIO(file_path_or_bytes) if isinstance(file_path_or_bytes, bytes) else file_path_or_bytes
        wb = load_workbook(buf, read_only=True, data_only=True)
    else:
        wb = load_workbook(file_path_or_bytes, read_only=True, data_only=True)

    try:
        return parse_xlsx_workbook(wb, estado, referencia)
    finally:
        wb.close()


def _detect_estado_from_filename(filename: str) -> Optional[str]:
    """Tenta detectar o estado a partir do nome do arquivo."""
    name_upper = _normalise(filename)
    # Tenta encontrar a UF no nome do arquivo (ex: "SINAPI_ref_Preco_SP_...")
    for uf in UF_NAMES:
        # Match UF isolated (not part of a word)
        if re.search(rf"(?<![A-Z]){uf}(?![A-Z])", name_upper):
            return uf
    return None


def _detect_referencia_from_filename(filename: str) -> Optional[str]:
    """Tenta detectar a referência (YYYY-MM) do nome do arquivo/ZIP."""
    # Patterns: "2026-01", "202601", "2026_01"
    m = re.search(r"(20\d{2})[-_]?(0[1-9]|1[0-2])", filename)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    return None


def load_zip_file(file_path_or_bytes) -> dict:
    """Carrega dados de um arquivo ZIP SINAPI (contendo XLSX por estado).

    Suporta dois formatos:
    - Formato simples: XLSX com múltiplas abas diretamente no ZIP.
    - Formato real Caixa: Pastas por UF/regime com arquivos XLSX separados
      (Preco_Ref_Insumos, Composicoes_Sintetico, Composicoes_Analitico).

    Args:
        file_path_or_bytes: Caminho do arquivo ZIP ou bytes.

    Returns:
        dict com dados combinados de todos os estados/arquivos.
    """
    if isinstance(file_path_or_bytes, (bytes, io.BytesIO)):
        buf = io.BytesIO(file_path_or_bytes) if isinstance(file_path_or_bytes, bytes) else file_path_or_bytes
        zf = zipfile.ZipFile(buf)
    else:
        zf = zipfile.ZipFile(file_path_or_bytes)

    # Detectar referência do nome do ZIP
    zip_name = getattr(file_path_or_bytes, "name", str(file_path_or_bytes))
    referencia = _detect_referencia_from_filename(str(zip_name)) or "2026-01"

    combined = {"insumos": [], "composicoes": [], "analitico": []}

    try:
        for name in zf.namelist():
            if not name.lower().endswith(".xlsx"):
                continue
            if name.startswith("__MACOSX") or name.startswith("."):
                continue

            estado = _detect_estado_from_filename(name)
            if not estado:
                continue

            ref = _detect_referencia_from_filename(name) or referencia

            # Detect file type from filename (real SINAPI format)
            file_type = _detect_file_type_from_filename(name)

            xlsx_bytes = zf.read(name)

            if file_type is not None:
                # Real format: regime from folder/filename, type from filename
                regime = _detect_regime(name)
                wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
                try:
                    data = parse_single_type_xlsx(wb, file_type, estado, ref, regime)
                finally:
                    wb.close()
            else:
                # Legacy format: multi-sheet XLSX
                data = load_xlsx_file(io.BytesIO(xlsx_bytes), estado, ref)

            combined["insumos"].extend(data["insumos"])
            combined["composicoes"].extend(data["composicoes"])
            combined["analitico"].extend(data["analitico"])
    finally:
        zf.close()

    return combined
